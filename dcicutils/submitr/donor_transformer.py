"""Analysis and selective transformation of ProtectedDonor workbook references.

The public entry points are :func:`analyze_protected_donors` and
:class:`ProtectedDonorWorkbookTransformer`.  ``analyze`` inspects non-empty
rows in the five protected-item sheets and classifies each ``donor`` value as
``plain_donor``, ``protected_donor``, ``missing``, or ``invalid``.  ``transform``
rewrites only referenced plain Donors and adds the corresponding
ProtectedDonor rows.

The submitr worker should enable this in its ``CustomExcel`` class with
``transform_protected_donor=True`` and pass its portal object through
``CustomExcel.with_portal``.  Portal lookups are deliberately made here,
before StructuredDataSet reads the workbook; ``structured_data.py`` does not
need to know about this SMaHT-specific conversion.
"""

from dataclasses import dataclass
from enum import Enum
from typing import Callable, Dict, Iterable, List, Optional, Set

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet


DONOR_SHEET = "Donor"
PROTECTED_DONOR_SHEET = "ProtectedDonor"
SUBMITTED_ID_COLUMN = "submitted_id"
PROTECTED_DONOR_COLUMN = "protected_donor"
DONOR_LINK_COLUMN = "donor"
DONOR_TOKEN = "_DONOR_"
PROTECTED_DONOR_TOKEN = "_PROTECTED-DONOR_"
PROTECTED_DONOR_STATUS = "in review"
PROTECTED_REFERENCE_SHEETS = frozenset({
    "Demographic",
    "DeathCircumstances",
    "FamilyHistory",
    "MedicalHistory",
    "TissueCollection",
})
SERVER_MANAGED_COLUMNS = frozenset({"accession", "uuid", "alternate_accessions"})


class ProtectedDonorTransformError(ValueError):
    """Raised when a workbook cannot be safely transformed."""


class DonorReferenceKind(str, Enum):
    PLAIN_DONOR = "plain_donor"
    PROTECTED_DONOR = "protected_donor"
    MISSING = "missing"
    INVALID = "invalid"


@dataclass(frozen=True)
class ProtectedDonorAnalysis:
    """The result of analyzing actual protected-item data rows.

    ``references`` maps each distinct non-empty donor reference to its kind.
    ``plain_donor_ids`` and ``protected_donor_ids`` contain references found
    in the workbook (or, for ProtectedDonor values, in the portal).  A value
    in ``missing_references`` has the right identifier shape but was not found
    where it can be resolved.  ``invalid_references`` has no valid Donor or
    ProtectedDonor identifier shape.
    """

    references: Dict[str, DonorReferenceKind]
    plain_donor_ids: frozenset
    protected_donor_ids: frozenset
    missing_references: frozenset
    invalid_references: frozenset
    missing_plain_donor_ids: frozenset
    missing_protected_donor_ids: frozenset

    @property
    def referenced_plain_donor_ids(self) -> frozenset:
        return self.plain_donor_ids

    @property
    def referenced_protected_donor_ids(self) -> frozenset:
        return self.protected_donor_ids

    @property
    def needs_transformation(self) -> bool:
        return bool(self.plain_donor_ids)


def to_protected_donor_submitted_id(submitted_id: str) -> str:
    """Return the ProtectedDonor ID corresponding to a plain Donor ID."""
    if not isinstance(submitted_id, str) or DONOR_TOKEN not in submitted_id:
        raise ProtectedDonorTransformError(
            f"Cannot derive ProtectedDonor submitted_id from {submitted_id!r}; "
            f"expected token {DONOR_TOKEN!r}."
        )
    return submitted_id.replace(DONOR_TOKEN, PROTECTED_DONOR_TOKEN, 1)


def analyze_protected_donors(workbook: openpyxl.Workbook, portal=None,
                             effective_sheet_name: Optional[Callable[[str], str]] = None
                             ) -> ProtectedDonorAnalysis:
    """Analyze protected-item donor references in ``workbook``.

    Only rows with data in one of :data:`PROTECTED_REFERENCE_SHEETS` count.
    A ProtectedDonor ID not present in the workbook is checked through
    ``portal`` when supplied.  The portal is expected to provide either
    ``ref_exists(type_name, value)`` (the StructuredDataSet portal wrapper)
    or ``get_metadata('/ProtectedDonor/<submitted_id>')`` (portal clients).
    """
    transformer = ProtectedDonorWorkbookTransformer(effective_sheet_name=effective_sheet_name)
    return transformer.analyze(workbook, portal=portal)


class ProtectedDonorWorkbookTransformer:
    """Selectively transform workbook references from Donor to ProtectedDonor."""

    def __init__(self, effective_sheet_name: Optional[Callable[[str], str]] = None,
                 portal=None) -> None:
        self._effective_sheet_name = effective_sheet_name or (lambda sheet_name: sheet_name)
        self._portal = portal

    def analyze(self, workbook: openpyxl.Workbook, portal=None) -> ProtectedDonorAnalysis:
        """Classify references without changing ``workbook``."""
        portal = self._portal if portal is None else portal
        donor_ids = self._submitted_ids(self._sheets_by_effective_name(workbook, DONOR_SHEET))
        protected_ids = self._submitted_ids(
            self._sheets_by_effective_name(workbook, PROTECTED_DONOR_SHEET)
        )
        references = self._referenced_donor_values(workbook)
        classifications = {}
        plain_ids = set()
        protected_found = set()
        missing_plain = set()
        missing_protected = set()
        invalid = set()

        for reference in references:
            if DONOR_TOKEN in reference:
                if reference in donor_ids:
                    classifications[reference] = DonorReferenceKind.PLAIN_DONOR
                    plain_ids.add(reference)
                else:
                    classifications[reference] = DonorReferenceKind.MISSING
                    missing_plain.add(reference)
            elif PROTECTED_DONOR_TOKEN in reference:
                if reference in protected_ids or self._protected_donor_exists_in_portal(portal, reference):
                    classifications[reference] = DonorReferenceKind.PROTECTED_DONOR
                    protected_found.add(reference)
                else:
                    classifications[reference] = DonorReferenceKind.MISSING
                    missing_protected.add(reference)
            else:
                classifications[reference] = DonorReferenceKind.INVALID
                invalid.add(reference)

        return ProtectedDonorAnalysis(
            references=classifications,
            plain_donor_ids=frozenset(plain_ids),
            protected_donor_ids=frozenset(protected_found),
            missing_references=frozenset(missing_plain | missing_protected),
            invalid_references=frozenset(invalid),
            missing_plain_donor_ids=frozenset(missing_plain),
            missing_protected_donor_ids=frozenset(missing_protected),
        )

    def transform(self, workbook: openpyxl.Workbook, portal=None) -> bool:
        """Apply the analyzed selective conversion in place.

        Missing or invalid references fail before any cell is changed.  Missing
        plain Donors are always errors even when a portal is available; a
        ProtectedDonor can instead be resolved by that portal.
        """
        analysis = self.analyze(workbook, portal=portal)
        if analysis.invalid_references:
            raise ProtectedDonorTransformError(
                "Invalid donor reference(s) in protected-item data rows: "
                + self._format_ids(analysis.invalid_references)
            )
        if analysis.missing_plain_donor_ids:
            raise ProtectedDonorTransformError(
                "Plain Donor reference(s) are absent from the workbook Donor sheet: "
                + self._format_ids(analysis.missing_plain_donor_ids)
            )
        if analysis.missing_protected_donor_ids:
            raise ProtectedDonorTransformError(
                "ProtectedDonor reference(s) are absent from the workbook and portal: "
                + self._format_ids(analysis.missing_protected_donor_ids)
            )
        if not analysis.plain_donor_ids:
            return False

        donor_sheets = self._sheets_by_effective_name(workbook, DONOR_SHEET)
        protected_sheets = self._sheets_by_effective_name(workbook, PROTECTED_DONOR_SHEET)
        workbook_protected_ids = self._submitted_ids(protected_sheets)
        rows_to_add = []

        # Validate all existing protected_donor cells before mutating any row.
        for sheet in donor_sheets:
            headers = self._headers(sheet)
            submitted_id_column = self._column_index(headers, SUBMITTED_ID_COLUMN)
            protected_column = self._column_index(headers, PROTECTED_DONOR_COLUMN)
            if submitted_id_column is None:
                continue
            for row_number in range(2, sheet.max_row + 1):
                donor_id = self._normalized(sheet.cell(row_number, submitted_id_column).value)
                if donor_id not in analysis.plain_donor_ids:
                    continue
                expected = to_protected_donor_submitted_id(donor_id)
                if protected_column is not None:
                    existing = self._normalized(sheet.cell(row_number, protected_column).value)
                    if existing and existing != expected:
                        raise ProtectedDonorTransformError(
                            f"Donor row {row_number} protected_donor value "
                            f"{sheet.cell(row_number, protected_column).value!r} does not match "
                            f"expected {expected!r}."
                        )
                if expected not in workbook_protected_ids:
                    rows_to_add.append((sheet, row_number, headers, expected))

        if not rows_to_add:
            # References may already have workbook ProtectedDonor rows.  The
            # donor-link rewrite is still useful only for newly plain links.
            self._add_protected_donor_columns(workbook, analysis)
            self._rewrite_donor_links(workbook, analysis)
            return True

        protected_sheet = self._protected_sheet_for_append(workbook, protected_sheets, donor_sheets[0])
        for sheet, row_number, headers, protected_id in rows_to_add:
            self._append_protected_row(protected_sheet, self._protected_row_from_donor_row(
                sheet, row_number, headers, protected_id
            ))
            workbook_protected_ids.add(protected_id)

        self._add_protected_donor_columns(workbook, analysis)
        self._rewrite_donor_links(workbook, analysis)
        return True

    def _add_protected_donor_columns(self, workbook, analysis: ProtectedDonorAnalysis) -> None:
        for sheet in self._sheets_by_effective_name(workbook, DONOR_SHEET):
            headers = self._headers(sheet)
            submitted_id_column = self._column_index(headers, SUBMITTED_ID_COLUMN)
            if submitted_id_column is None:
                continue
            protected_column = self._ensure_column(sheet, headers, PROTECTED_DONOR_COLUMN)
            for row_number in range(2, sheet.max_row + 1):
                donor_id = self._normalized(sheet.cell(row_number, submitted_id_column).value)
                if donor_id in analysis.plain_donor_ids:
                    sheet.cell(row_number, protected_column).value = to_protected_donor_submitted_id(donor_id)

    def _rewrite_donor_links(self, workbook, analysis: ProtectedDonorAnalysis) -> None:
        for sheet in self._protected_reference_sheets(workbook):
            headers = self._headers(sheet)
            donor_column = self._column_index(headers, DONOR_LINK_COLUMN)
            if donor_column is None:
                continue
            for row_number in range(2, sheet.max_row + 1):
                value = self._normalized(sheet.cell(row_number, donor_column).value)
                if value in analysis.plain_donor_ids:
                    sheet.cell(row_number, donor_column).value = to_protected_donor_submitted_id(value)

    def _protected_sheet_for_append(self, workbook, protected_sheets, donor_sheet):
        if protected_sheets:
            return protected_sheets[0]
        name = PROTECTED_DONOR_SHEET
        donor_title = donor_sheet.title
        prefix, separator, _ = donor_title.rpartition("_")
        if separator and prefix:
            name = f"{prefix}_{PROTECTED_DONOR_SHEET}"
        return workbook.create_sheet(name)

    def _protected_row_from_donor_row(self, sheet, row_number, headers, protected_id):
        row = {}
        for column_number, header in enumerate(headers, start=1):
            if not header or header in SERVER_MANAGED_COLUMNS or header == PROTECTED_DONOR_COLUMN:
                continue
            value = sheet.cell(row_number, column_number).value
            if header == SUBMITTED_ID_COLUMN:
                value = protected_id
            elif header == "status":
                value = PROTECTED_DONOR_STATUS
            row[header] = value
        row.setdefault("status", PROTECTED_DONOR_STATUS)
        return row

    @staticmethod
    def _append_protected_row(sheet, row):
        headers = ProtectedDonorWorkbookTransformer._headers(sheet)
        if not headers:
            headers = list(row.keys())
            for column_number, header in enumerate(headers, start=1):
                sheet.cell(1, column_number).value = header
        row_number = sheet.max_row + 1 if sheet.max_row else 2
        for column_number, header in enumerate(headers, start=1):
            sheet.cell(row_number, column_number).value = row.get(header)

    @staticmethod
    def _collect_referenced_donor_values(workbook, effective_sheet_name=None) -> Set[str]:
        effective_sheet_name = effective_sheet_name or (lambda name: name)
        result = set()
        for sheet in workbook.worksheets:
            if ProtectedDonorWorkbookTransformer._is_hidden_sheet(sheet) or \
                    effective_sheet_name(sheet.title) not in PROTECTED_REFERENCE_SHEETS:
                continue
            headers = ProtectedDonorWorkbookTransformer._headers(sheet)
            donor_column = ProtectedDonorWorkbookTransformer._column_index(headers, DONOR_LINK_COLUMN)
            if donor_column is None:
                continue
            for row_number in range(2, sheet.max_row + 1):
                value = ProtectedDonorWorkbookTransformer._normalized(
                    sheet.cell(row_number, donor_column).value
                )
                if value:
                    result.add(value)
        return result

    def _protected_reference_sheets(self, workbook):
        return [sheet for sheet in self._visible_sheets(workbook)
                if self._effective_sheet_name(sheet.title) in PROTECTED_REFERENCE_SHEETS]

    def _sheets_by_effective_name(self, workbook, name):
        return [sheet for sheet in self._visible_sheets(workbook)
                if self._effective_sheet_name(sheet.title) == name]

    @staticmethod
    def _visible_sheets(workbook):
        return [sheet for sheet in workbook.worksheets
                if not ProtectedDonorWorkbookTransformer._is_hidden_sheet(sheet)]

    @staticmethod
    def _rows_by_submitted_id(sheets):
        result = {}
        for sheet in sheets:
            headers = ProtectedDonorWorkbookTransformer._headers(sheet)
            column = ProtectedDonorWorkbookTransformer._column_index(headers, SUBMITTED_ID_COLUMN)
            if column is None:
                continue
            for row_number in range(2, sheet.max_row + 1):
                value = ProtectedDonorWorkbookTransformer._normalized(sheet.cell(row_number, column).value)
                if value:
                    result[value] = (sheet, row_number, headers)
        return result

    @staticmethod
    def _submitted_ids(sheets):
        return set(ProtectedDonorWorkbookTransformer._rows_by_submitted_id(sheets))

    @staticmethod
    def _headers(sheet) -> List[str]:
        headers = []
        for cell in sheet[1]:
            if cell.value is None or not str(cell.value).strip():
                break
            headers.append(str(cell.value).strip())
        return headers

    @staticmethod
    def _column_index(headers: List[str], column_name: str) -> Optional[int]:
        try:
            return headers.index(column_name) + 1
        except ValueError:
            return None

    @staticmethod
    def _ensure_column(sheet, headers: List[str], column_name: str) -> int:
        column = ProtectedDonorWorkbookTransformer._column_index(headers, column_name)
        if column is None:
            column = len(headers) + 1
            sheet.cell(1, column).value = column_name
            headers.append(column_name)
        return column

    @staticmethod
    def _normalized(value) -> str:
        return str(value).strip() if value is not None else ""

    @staticmethod
    def _is_hidden_sheet(sheet: Worksheet) -> bool:
        title = sheet.title
        return sheet.sheet_state == "hidden" or any(
            title.startswith(left) and title.endswith(right)
            for left, right in (("(", ")"), ("[", "]"), ("{", "}"), ("<", ">"))
        )

    @staticmethod
    def _format_ids(values: Iterable[str]) -> str:
        return ", ".join(repr(value) for value in sorted(values))

    @staticmethod
    def _protected_donor_exists_in_portal(portal, submitted_id: str) -> bool:
        if portal is None:
            return False
        ref_exists = getattr(portal, "ref_exists", None)
        if callable(ref_exists):
            for args in ((PROTECTED_DONOR_SHEET, submitted_id),
                         (f"/{PROTECTED_DONOR_SHEET}/{submitted_id}",)):
                try:
                    if ref_exists(*args):
                        return True
                except (AttributeError, TypeError, ValueError):
                    continue
        get_metadata = getattr(portal, "get_metadata", None)
        if callable(get_metadata):
            path = f"/{PROTECTED_DONOR_SHEET}/{submitted_id}"
            for kwargs in ({"raw": True}, {}):
                try:
                    result = get_metadata(path, **kwargs)
                    if result:
                        return True
                    break
                except (AttributeError, TypeError, ValueError, KeyError):
                    continue
                except Exception as error:
                    if "notfound" in str(error).lower() or "404" in str(error):
                        break
        return False

    def _referenced_donor_values(self, workbook):
        return self._collect_referenced_donor_values(
            workbook, effective_sheet_name=self._effective_sheet_name
        )
