import openpyxl
import pytest

from dcicutils.structured_data import StructuredDataSet
from dcicutils.submitr.custom_excel import CustomExcel
from dcicutils.submitr.donor_transformer import (
    DonorReferenceKind,
    ProtectedDonorLookupError,
    ProtectedDonorTransformError,
    ProtectedDonorWorkbookTransformer,
    analyze_protected_donors,
    lookup_protected_donor,
)


def _workbook(sheets):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    return workbook


def test_analysis_uses_only_actual_rows_and_the_five_protected_sheets():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["A_DONOR_1"], [None]],
        "DeathCircumstances": [["donor"], ["A_PROTECTED-DONOR_2"]],
        "FamilyHistory": [["donor"], ["not-an-id"]],
        "MedicalHistory": [["donor"], ["A_DONOR_MISSING"]],
        "TissueCollection": [["donor"], ["A_DONOR_1"]],
        "Tissue": [["donor"], ["A_DONOR_IGNORED"]],
    })

    analysis = analyze_protected_donors(workbook)

    assert analysis.references["A_DONOR_1"] == DonorReferenceKind.PLAIN_DONOR
    assert analysis.references["A_PROTECTED-DONOR_2"] == DonorReferenceKind.MISSING
    assert analysis.references["not-an-id"] == DonorReferenceKind.INVALID
    assert analysis.missing_plain_donor_ids == frozenset({"A_DONOR_MISSING"})
    assert analysis.invalid_references == frozenset({"not-an-id"})
    assert "A_DONOR_IGNORED" not in analysis.references


def test_transform_is_selective_and_supports_mixed_references():
    workbook = _workbook({
        "Donor": [
            ["submitted_id", "external_id", "status"],
            ["A_DONOR_REFERENCED", "one", "released"],
            ["A_DONOR_UNREFERENCED", "two", "released"],
        ],
        "ProtectedDonor": [
            ["submitted_id", "external_id", "status"],
            ["A_PROTECTED-DONOR_EXISTING", "existing", "in review"],
        ],
        "Demographic": [["donor"], ["A_DONOR_REFERENCED"]],
        "MedicalHistory": [["donor"], ["A_PROTECTED-DONOR_EXISTING"]],
        "Tissue": [["donor"], ["A_DONOR_UNREFERENCED"]],
    })

    assert ProtectedDonorWorkbookTransformer().transform(workbook) is True

    donors = workbook["Donor"]
    assert donors.cell(1, 4).value == "protected_donor"
    assert donors.cell(2, 4).value == "A_PROTECTED-DONOR_REFERENCED"
    assert donors.cell(3, 4).value is None
    protected = workbook["ProtectedDonor"]
    assert protected.max_row == 3
    assert protected.cell(3, 1).value == "A_PROTECTED-DONOR_REFERENCED"
    assert protected.cell(3, 2).value == "one"
    assert protected.cell(3, 3).value == "in review"
    assert workbook["Demographic"].cell(2, 1).value == "A_PROTECTED-DONOR_REFERENCED"
    assert workbook["MedicalHistory"].cell(2, 1).value == "A_PROTECTED-DONOR_EXISTING"
    assert workbook["Tissue"].cell(2, 1).value == "A_DONOR_UNREFERENCED"


def test_invalid_donor_reference_is_a_clear_transform_error():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_PRESENT"]],
        "Demographic": [["donor"], ["not-a-donor-id"]],
    })

    with pytest.raises(ProtectedDonorTransformError, match="Invalid donor reference"):
        ProtectedDonorWorkbookTransformer().transform(workbook)


def test_missing_plain_donor_is_a_clear_transform_error():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_PRESENT"]],
        "Demographic": [["donor"], ["A_DONOR_ABSENT"]],
    })

    with pytest.raises(ProtectedDonorTransformError, match="absent from the workbook Donor sheet"):
        ProtectedDonorWorkbookTransformer().transform(workbook)


def test_portal_only_protected_donor_is_accepted():
    class Portal:
        def get_metadata(self, path, **kwargs):
            assert path == "/ProtectedDonor/A_PROTECTED-DONOR_PORTAL"
            return {"uuid": "portal-uuid"}

    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_UNREFERENCED"]],
        "Demographic": [["donor"], ["A_PROTECTED-DONOR_PORTAL"]],
    })

    analysis = ProtectedDonorWorkbookTransformer(portal=Portal()).analyze(workbook)
    assert analysis.protected_donor_ids == frozenset({"A_PROTECTED-DONOR_PORTAL"})
    assert ProtectedDonorWorkbookTransformer(portal=Portal()).transform(workbook) is False
    assert "ProtectedDonor" not in workbook.sheetnames


def test_existing_protected_sheet_is_extended_incrementally():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_NEW"]],
        "ProtectedDonor": [["submitted_id", "status"], ["A_PROTECTED-DONOR_OLD", "in review"]],
        "FamilyHistory": [["donor"], ["A_DONOR_NEW"]],
    })

    ProtectedDonorWorkbookTransformer().transform(workbook)

    assert workbook["ProtectedDonor"].max_row == 3
    assert workbook["ProtectedDonor"].cell(2, 1).value == "A_PROTECTED-DONOR_OLD"
    assert workbook["ProtectedDonor"].cell(3, 1).value == "A_PROTECTED-DONOR_NEW"


def test_custom_excel_opt_in_integrates_transform_before_structured_data(tmp_path):
    path = tmp_path / "donor.xlsx"
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["A_DONOR_1"]],
    })
    workbook.save(path)

    data = StructuredDataSet(
        file=str(path),
        portal=None,
        excel_class=CustomExcel.with_portal(None, transform_protected_donor=True),
    ).data

    assert data["ProtectedDonor"][0]["submitted_id"] == "A_PROTECTED-DONOR_1"
    assert data["Demographic"][0]["donor"] == "A_PROTECTED-DONOR_1"


def test_prefixed_protected_item_sheet_uses_canonical_name():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "DSA_Demographic": [["donor"], ["A_DONOR_1"]],
    })

    transformer = ProtectedDonorWorkbookTransformer(
        effective_sheet_name=lambda name: name.split("_", 1)[-1]
    )
    assert transformer.transform(workbook) is True
    assert workbook["DSA_Demographic"]["A2"].value == "A_PROTECTED-DONOR_1"


def test_accession_lookup_uses_portal_root_path():
    portal = _TypedPortal({"SMAPD1234567": {"uuid": "portal-uuid", "@type": ["ProtectedDonor", "Item"]}})
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["SMAPD1234567"]],
    })

    analysis = ProtectedDonorWorkbookTransformer(portal=portal).analyze(workbook)
    assert analysis.references["SMAPD1234567"] == DonorReferenceKind.PROTECTED_DONOR
    assert portal.paths == ["/SMAPD1234567"]


@pytest.mark.parametrize("item", [{"uuid": "portal-uuid", "@type": ["Donor", "AbstractDonor", "Item"]},
                                  {"uuid": "portal-uuid"}])
def test_accession_of_a_non_protected_donor_item_is_not_a_protected_donor(item):
    portal = _TypedPortal({"SMADO1234567": item})
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["SMADO1234567"]],
    })

    assert ProtectedDonorWorkbookTransformer(portal=portal).analyze(workbook).references["SMADO1234567"] == \
        DonorReferenceKind.INVALID


@pytest.mark.parametrize("error,kind", [
    (Exception("HTTPForbidden: /SMAPD1234567"), DonorReferenceKind.PERMISSION_DENIED),
    (Exception("Bad response: 500 Internal Server Error"), DonorReferenceKind.LOOKUP_FAILED),
])
def test_accession_lookup_failures_are_distinct_from_absence(error, kind):
    portal = _TypedPortal(error=error)
    workbook = _workbook({"Demographic": [["donor"], ["SMAPD1234567"]]})

    assert ProtectedDonorWorkbookTransformer(portal=portal).analyze(workbook).references["SMAPD1234567"] == kind
    with pytest.raises(ProtectedDonorLookupError, match="not known to be absent"):
        ProtectedDonorWorkbookTransformer(portal=portal).transform(workbook)


@pytest.mark.parametrize("title", ["demographic", "DSA_demographic", "Death Circumstances"])
def test_custom_excel_matches_protected_item_sheets_as_the_reader_names_them(tmp_path, title):
    path = tmp_path / "donor.xlsx"
    _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        title: [["donor"], ["A_DONOR_1"]],
    }).save(path)

    data = StructuredDataSet(
        file=str(path),
        portal=None,
        excel_class=CustomExcel.with_portal(None, transform_protected_donor=True),
    ).data

    assert data["ProtectedDonor"][0]["submitted_id"] == "A_PROTECTED-DONOR_1"
    linked = [items[0]["donor"] for name, items in data.items() if name not in ("Donor", "ProtectedDonor")]
    assert linked == ["A_PROTECTED-DONOR_1"]


def test_hidden_and_unlisted_sheets_are_not_transformed():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "(Demographic)": [["donor"], ["A_DONOR_1"]],
        "Tissue": [["donor"], ["A_DONOR_1"]],
    })
    workbook["(Demographic)"].sheet_state = "hidden"

    assert ProtectedDonorWorkbookTransformer().transform(workbook) is False
    assert "ProtectedDonor" not in workbook.sheetnames


# --- Regression tests for PR review findings ---

def _formatted_empty_rows(sheet, count):
    # Formatted-but-empty cells extend max_row without the reader ever seeing data there.
    for offset in range(1, count + 1):
        sheet.cell(sheet.max_row + 1, 1).number_format = "@"
    return sheet


def _read_rows(path, sheet_name):
    return StructuredDataSet(file=str(path), portal=None, norefs=True).data.get(sheet_name, [])


def test_existing_protected_sheet_gains_headers_for_copied_fields():
    workbook = _workbook({
        "Donor": [["submitted_id", "external_id", "status"], ["A_DONOR_NEW", "ext-1", "released"]],
        "ProtectedDonor": [["submitted_id"], ["A_PROTECTED-DONOR_OLD"]],
        "FamilyHistory": [["donor"], ["A_DONOR_NEW"]],
    })

    ProtectedDonorWorkbookTransformer().transform(workbook)

    protected = workbook["ProtectedDonor"]
    headers = [cell.value for cell in protected[1]]
    assert headers == ["submitted_id", "external_id", "status"]
    assert [cell.value for cell in protected[2]] == ["A_PROTECTED-DONOR_OLD", None, None]
    assert [cell.value for cell in protected[3]] == ["A_PROTECTED-DONOR_NEW", "ext-1", "in review"]


def test_existing_protected_sheet_with_unsafe_extra_columns_is_rejected_unchanged():
    workbook = _workbook({
        "Donor": [["submitted_id", "external_id"], ["A_DONOR_NEW", "ext-1"]],
        "ProtectedDonor": [["submitted_id", None, "orphan"], ["A_PROTECTED-DONOR_OLD", None, "stale"]],
        "FamilyHistory": [["donor"], ["A_DONOR_NEW"]],
    })

    with pytest.raises(ProtectedDonorTransformError, match="columns cannot be added"):
        ProtectedDonorWorkbookTransformer().transform(workbook)

    assert workbook["ProtectedDonor"].max_row == 2
    assert workbook["Donor"].max_column == 2
    assert workbook["FamilyHistory"]["A2"].value == "A_DONOR_NEW"


def test_generated_rows_are_appended_before_formatted_empty_rows_and_stay_visible(tmp_path):
    workbook = _workbook({
        "Donor": [["submitted_id", "status"], ["A_DONOR_1", "released"], ["A_DONOR_2", "released"]],
        "ProtectedDonor": [["submitted_id", "status"], ["A_PROTECTED-DONOR_OLD", "in review"]],
        "FamilyHistory": [["donor"], ["A_DONOR_1"], ["A_DONOR_2"]],
    })
    _formatted_empty_rows(workbook["ProtectedDonor"], 5)
    assert workbook["ProtectedDonor"].max_row > 6

    ProtectedDonorWorkbookTransformer().transform(workbook)

    path = tmp_path / "out.xlsx"
    workbook.save(path)
    ids = [row["submitted_id"] for row in _read_rows(path, "ProtectedDonor")]
    assert ids == ["A_PROTECTED-DONOR_OLD", "A_PROTECTED-DONOR_1", "A_PROTECTED-DONOR_2"]


def test_content_after_the_empty_row_terminator_blocks_append_and_leaves_workbook_unchanged():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "ProtectedDonor": [["submitted_id"], ["A_PROTECTED-DONOR_OLD"], [None], ["A_PROTECTED-DONOR_STALE"]],
        "FamilyHistory": [["donor"], ["A_DONOR_1"]],
    })

    with pytest.raises(ProtectedDonorTransformError, match="after its first empty row"):
        ProtectedDonorWorkbookTransformer().transform(workbook)

    assert workbook["ProtectedDonor"].max_row == 4
    assert workbook["ProtectedDonor"]["A3"].value is None
    assert workbook["FamilyHistory"]["A2"].value == "A_DONOR_1"
    assert "protected_donor" not in [cell.value for cell in workbook["Donor"][1]]


def test_references_after_the_empty_row_terminator_are_not_scanned():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"], [None], ["A_DONOR_STALE"]],
        "Demographic": [["donor"], ["A_DONOR_1"], [None], ["not-an-id"], ["A_DONOR_STALE"]],
    })

    analysis = analyze_protected_donors(workbook)
    assert set(analysis.references) == {"A_DONOR_1"}

    assert ProtectedDonorWorkbookTransformer().transform(workbook) is True
    assert workbook["Demographic"]["A2"].value == "A_PROTECTED-DONOR_1"
    # Rows after the terminator are never ingested, so they are never rewritten either.
    assert workbook["Demographic"]["A4"].value == "not-an-id"
    assert workbook["Demographic"]["A5"].value == "A_DONOR_STALE"
    assert workbook["Donor"].cell(4, 1).value == "A_DONOR_STALE"
    assert workbook["Donor"].max_column == 2
    assert workbook["Donor"].cell(4, 2).value is None


class _TypedPortal:
    """Portal stub answering typed ProtectedDonor lookups by identifier."""

    def __init__(self, items=None, error=None):
        self.items = items or {}
        self.error = error
        self.paths = []

    def get_metadata(self, path, **kwargs):
        self.paths.append(path)
        if self.error:
            raise self.error
        prefix = "/ProtectedDonor/" if path.startswith("/ProtectedDonor/") else "/"
        if path[len(prefix):] in self.items:
            return self.items[path[len(prefix):]]
        raise Exception(f"Bad status code for GET request for https://portal{path}: 404. Reason: Not Found")


UUID = "11111111-2222-3333-4444-555555555555"
ACCESSION = "SMAPD1234567"


@pytest.mark.parametrize("identifier", [UUID, ACCESSION])
def test_uuid_and_accession_references_to_existing_protected_donors_are_accepted(identifier):
    portal = _TypedPortal({identifier: {"uuid": UUID, "@type": ["ProtectedDonor", "Item"]}})
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], [identifier], ["A_DONOR_1"]],
    })

    analysis = ProtectedDonorWorkbookTransformer(portal=portal).analyze(workbook)
    assert analysis.references[identifier] == DonorReferenceKind.PROTECTED_DONOR
    assert not analysis.invalid_references

    assert ProtectedDonorWorkbookTransformer(portal=portal).transform(workbook) is True
    assert workbook["Demographic"]["A2"].value == identifier
    assert workbook["Demographic"]["A3"].value == "A_PROTECTED-DONOR_1"


def test_workbook_protected_donor_uuid_and_accession_are_resolved_without_the_portal():
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_UNREFERENCED"]],
        "ProtectedDonor": [["submitted_id", "uuid", "accession"], ["A_PROTECTED-DONOR_X", UUID, ACCESSION]],
        "Demographic": [["donor"], [UUID], [ACCESSION]],
    })

    analysis = analyze_protected_donors(workbook)

    assert analysis.protected_donor_ids == frozenset({UUID, ACCESSION})


def test_identifier_resolving_to_a_different_type_is_not_a_protected_donor():
    portal = _TypedPortal({UUID: {"uuid": UUID, "@type": ["Donor", "Item"]}})
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], [UUID]],
    })

    analysis = ProtectedDonorWorkbookTransformer(portal=portal).analyze(workbook)

    assert analysis.references[UUID] == DonorReferenceKind.INVALID
    with pytest.raises(ProtectedDonorTransformError, match="Invalid donor reference"):
        ProtectedDonorWorkbookTransformer(portal=portal).transform(workbook)


def test_unresolvable_identifier_without_portal_stays_invalid():
    workbook = _workbook({"Demographic": [["donor"], [UUID]]})

    assert analyze_protected_donors(workbook).references[UUID] == DonorReferenceKind.INVALID


@pytest.mark.parametrize("error,kind", [
    (Exception("Bad status code for GET request for https://p/x: 403. Reason: Forbidden"),
     DonorReferenceKind.PERMISSION_DENIED),
    (Exception("Bad status code for GET request for https://p/x: 401. Reason: Unauthorized"),
     DonorReferenceKind.PERMISSION_DENIED),
    (Exception("HTTPForbidden: no access"), DonorReferenceKind.PERMISSION_DENIED),
    (Exception("Bad status code for GET request for https://p/404-404: 503. Reason: Unavailable"),
     DonorReferenceKind.LOOKUP_FAILED),
    (ConnectionError("connection refused"), DonorReferenceKind.LOOKUP_FAILED),
    (TimeoutError("timed out"), DonorReferenceKind.LOOKUP_FAILED),
])
def test_lookup_failures_are_distinct_from_absence(error, kind):
    portal = _TypedPortal(error=error)
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["A_PROTECTED-DONOR_PORTAL"]],
    })

    analysis = ProtectedDonorWorkbookTransformer(portal=portal).analyze(workbook)

    assert analysis.references["A_PROTECTED-DONOR_PORTAL"] == kind
    assert not analysis.missing_references
    assert analysis.lookup_errors["A_PROTECTED-DONOR_PORTAL"]
    with pytest.raises(ProtectedDonorLookupError, match="not known to be absent") as caught:
        ProtectedDonorWorkbookTransformer(portal=portal).transform(workbook)
    assert "absent from the workbook and portal" not in str(caught.value)
    assert workbook["Demographic"]["A2"].value == "A_PROTECTED-DONOR_PORTAL"


@pytest.mark.parametrize("portal", [
    _TypedPortal(),
    _TypedPortal(error=Exception("HTTPNotFound: /ProtectedDonor/x")),
    _TypedPortal(items={"A_PROTECTED-DONOR_PORTAL": {"@type": ["HTTPNotFound", "Error"], "status": "error",
                                                     "code": 404}}),
])
def test_definitive_not_found_is_reported_as_absent(portal):
    workbook = _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["A_PROTECTED-DONOR_PORTAL"]],
    })

    analysis = ProtectedDonorWorkbookTransformer(portal=portal).analyze(workbook)

    assert analysis.references["A_PROTECTED-DONOR_PORTAL"] == DonorReferenceKind.MISSING
    with pytest.raises(ProtectedDonorTransformError, match="absent from the workbook and portal") as caught:
        ProtectedDonorWorkbookTransformer(portal=portal).transform(workbook)
    assert not isinstance(caught.value, ProtectedDonorLookupError)


def test_error_result_with_forbidden_type_is_a_permission_failure():
    portal = _TypedPortal({"A_PROTECTED-DONOR_PORTAL": {"@type": ["HTTPForbidden", "Error"], "status": "error"}})
    workbook = _workbook({"Demographic": [["donor"], ["A_PROTECTED-DONOR_PORTAL"]]})

    analysis = ProtectedDonorWorkbookTransformer(portal=portal).analyze(workbook)

    assert analysis.references["A_PROTECTED-DONOR_PORTAL"] == DonorReferenceKind.PERMISSION_DENIED


def test_lookup_identifier_is_path_quoted():
    portal = _TypedPortal({})
    lookup_protected_donor(portal, "a/b?c")

    assert portal.paths == ["/ProtectedDonor/a%2Fb%3Fc"]


def _counting_excel_class(path, **options):
    created = []

    class _Counting(CustomExcel):
        def __init__(self, *args, **kwargs):
            created.append(self)
            kwargs.setdefault("transform_protected_donor", True)
            kwargs.setdefault("transformed_workbook_path", str(path))
            super().__init__(*args, **kwargs)

    return _Counting, created


def test_progress_loading_does_not_save_during_the_counting_pass(tmp_path):
    input_path = tmp_path / "in.xlsx"
    _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["A_DONOR_1"]],
    }).save(input_path)
    staged = tmp_path / "staged.xlsx"
    excel_class, created = _counting_excel_class(staged)
    progress = []

    data = StructuredDataSet(file=str(input_path), portal=None, norefs=True,
                             excel_class=excel_class, progress=progress.append).data

    assert len(created) == 1
    assert progress
    assert data["Demographic"][0]["donor"] == "A_PROTECTED-DONOR_1"
    assert [row["submitted_id"] for row in openpyxl_rows(staged, "ProtectedDonor")] == ["A_PROTECTED-DONOR_1"]


def openpyxl_rows(path, sheet_name):
    sheet = openpyxl.load_workbook(path)[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    return [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def test_repeated_construction_still_cannot_clobber_the_staged_path(tmp_path):
    input_path = tmp_path / "in.xlsx"
    _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["A_DONOR_1"]],
    }).save(input_path)
    staged = tmp_path / "staged.xlsx"
    kwargs = dict(file=str(input_path), transform_protected_donor=True, transformed_workbook_path=str(staged))

    CustomExcel(**kwargs)
    original = staged.read_bytes()
    with pytest.raises(ValueError, match="already exists"):
        CustomExcel(**kwargs)

    assert staged.read_bytes() == original


def _write_donor_input(path):
    _workbook({
        "Donor": [["submitted_id"], ["A_DONOR_1"]],
        "Demographic": [["donor"], ["A_DONOR_1"]],
    }).save(path)


def test_owned_staging_path_is_replaced_when_existing_path_is_allowed(tmp_path):
    input_path = tmp_path / "in.xlsx"
    _write_donor_input(input_path)
    staged = tmp_path / "staged.xlsx"
    _workbook({"Stale": [["old"], ["content"]]}).save(staged)

    excel = CustomExcel(file=str(input_path), transform_protected_donor=True,
                        transformed_workbook_path=str(staged), allow_existing_staging_path=True)

    assert "ProtectedDonor" in excel.sheet_names
    assert "Stale" not in openpyxl.load_workbook(staged).sheetnames
    assert [row["submitted_id"] for row in openpyxl_rows(staged, "ProtectedDonor")] == ["A_PROTECTED-DONOR_1"]
    assert [row["donor"] for row in openpyxl_rows(staged, "Demographic")] == ["A_PROTECTED-DONOR_1"]
    assert sorted(p.name for p in tmp_path.iterdir()) == ["in.xlsx", "staged.xlsx"]


def test_failed_save_to_owned_staging_path_keeps_existing_file_and_leaves_no_temporary(tmp_path, monkeypatch):
    input_path = tmp_path / "in.xlsx"
    _write_donor_input(input_path)
    staged = tmp_path / "staged.xlsx"
    _workbook({"Stale": [["old"], ["content"]]}).save(staged)
    original = staged.read_bytes()

    def failing_save(self, filename):
        with open(filename, "wb") as stream:
            stream.write(b"partial")
        raise OSError("disk full")

    monkeypatch.setattr(openpyxl.Workbook, "save", failing_save)
    with pytest.raises(ValueError, match="Cannot save transformed workbook.*disk full"):
        CustomExcel(file=str(input_path), transform_protected_donor=True,
                    transformed_workbook_path=str(staged), allow_existing_staging_path=True)

    assert staged.read_bytes() == original
    assert sorted(p.name for p in tmp_path.iterdir()) == ["in.xlsx", "staged.xlsx"]


def test_construction_failure_does_not_raise_attribute_error_in_excel_del(tmp_path):
    import gc
    import sys

    unraisable = []
    previous_hook = sys.unraisablehook
    sys.unraisablehook = unraisable.append
    try:
        with pytest.raises(TypeError):
            CustomExcel(file=str(tmp_path / "in.xlsx"), unsupported_option=True)
        gc.collect()
    finally:
        sys.unraisablehook = previous_hook

    assert [repr(item.exc_value) for item in unraisable] == []
